"""
Exportación a Excel con el sistema de documentos QKT (Issue #373)
=================================================================
Contraparte de `core_erp/documentos.py` para hojas de cálculo: mismo nombre de
archivo (`QKT_<Tipo>_<…>.xlsx`), encabezado con la identidad del ERP y montos
como número (no como texto) para que se puedan sumar y filtrar.

Cada hoja es `(titulo, encabezados, filas)`; una celda `Decimal` sale con
formato de moneda y una fecha con formato `dd/mm/aaaa`.
"""
from datetime import date, datetime
from decimal import Decimal

from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from core_erp.documentos import EMPRESA

FORMATO_MONEDA = '"$"#,##0.00;-"$"#,##0.00'
FORMATO_FECHA = 'dd/mm/yyyy'
_ENCABEZADO = Font(bold=True, color='FFFFFF')
_FONDO = PatternFill('solid', fgColor='1B5E20')  # --verde-oscuro de documentos/_estilos.html


def libro_excel(titulo: str, hojas) -> Workbook:
    libro = Workbook()
    libro.remove(libro.active)
    libro.properties.title = titulo
    libro.properties.creator = EMPRESA
    for nombre, encabezados, filas in hojas:
        hoja = libro.create_sheet(nombre[:31])  # Excel no admite títulos de hoja más largos
        hoja.append(list(encabezados))
        for celda in hoja[1]:
            celda.font, celda.fill = _ENCABEZADO, _FONDO
        anchos = [len(str(e)) for e in encabezados]
        for fila in filas:
            hoja.append(list(fila))
            for i, valor in enumerate(fila):
                celda = hoja.cell(row=hoja.max_row, column=i + 1)
                if isinstance(valor, Decimal):
                    celda.number_format = FORMATO_MONEDA
                elif isinstance(valor, (date, datetime)):
                    celda.number_format = FORMATO_FECHA
                anchos[i] = max(anchos[i], len(str(valor)) if valor is not None else 0)
        for i, ancho in enumerate(anchos, start=1):
            hoja.column_dimensions[get_column_letter(i)].width = min(ancho + 2, 50)
        hoja.freeze_panes = 'A2'
    return libro


def respuesta_excel(titulo: str, hojas, nombre: str) -> HttpResponse:
    """`HttpResponse` de descarga con el libro armado por `libro_excel`."""
    respuesta = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    respuesta['Content-Disposition'] = f'attachment; filename="{nombre}"'
    libro_excel(titulo, hojas).save(respuesta)
    return respuesta
