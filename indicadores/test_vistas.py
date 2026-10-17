"""Tablero de Indicadores (Issue #402, fase 3): acceso, render y reglas del sistema de diseño."""
import re
from datetime import date
from decimal import Decimal
from pathlib import Path

from django.conf import settings
from django.contrib.auth.models import User
from django.test import TestCase
from django.urls import reverse

from core_erp.test_utils import login_superuser_con_totp
from indicadores.test_services import MesArmado

EMOJI = re.compile('[\U0001F300-\U0001FAFF☀-➿]')


D = Decimal


class AccesoTest(TestCase):

    @classmethod
    def setUpTestData(cls):
        cls.direccion = User.objects.create_superuser('direccion', 'd@x.mx', 'clave-de-prueba-123')
        cls.staff = User.objects.create_user('staff', 's@x.mx', 'clave-de-prueba-123', is_staff=True)

    def test_anonimo_va_al_login_y_staff_no_superusuario_recibe_403(self):
        url = reverse('indicadores:tablero')
        self.assertEqual(self.client.get(url).status_code, 302)
        self.client.force_login(self.staff)
        self.assertEqual(self.client.get(url).status_code, 403)

    def test_tablero_renderiza_con_base_vacia(self):
        login_superuser_con_totp(self.client, self.direccion)
        r = self.client.get(reverse('indicadores:tablero'))
        self.assertEqual(r.status_code, 200)
        self.assertContains(r, 'Este mes')
        self.assertContains(r, 'Revisa')           # avisos en vez de errores
        r = self.client.get(reverse('indicadores:tablero'), {'mes': 'basura'})
        self.assertEqual(r.status_code, 200)       # un mes inválido cae al mes en curso

    def test_el_menu_no_se_muestra_al_staff(self):
        self.client.force_login(self.staff)
        r = self.client.get(reverse('admin:index'))
        self.assertNotContains(r, 'indicadores/parametrosnegocio')


class TableroConDatosTest(MesArmado):

    def setUp(self):
        self.precios_de_febrero(guardar=True)
        login_superuser_con_totp(self.client, User.objects.create_superuser('dir', 'd@x.mx', 'clave-de-prueba-123'))

    def test_muestra_las_cifras_calculadas_a_mano(self):
        r = self.client.get(reverse('indicadores:tablero'), {'mes': '2026-02'})
        self.assertEqual(r.status_code, 200)
        # Febrero 2026 (MesArmado): venta neta 75,000 (los 74,000 de eventos, pasadías y hospedajes + 1,000 de arrendamiento);
        # margen ponderado 65 %, costos fijos 32,500 → utilidad 75,000 × 0.65 − 32,500 = 16,250. Equilibrio: 32,500 / 0.65 = 50,000
        # (ya cubierto) y 32,500 / 3,006.25 = 10.81 → 11 días; vendidos 18 de 28 vendibles = 64.3 %. Meta = venta del plan, 148,000.
        for esperado in ('$75,000.00', '$148,000.00', '$16,250.00', '18 de 11 días', 'Costos fijos cubiertos',
                         'Equilibrio alcanzado', '64.3 %', '28 días vendibles', 'Febrero 2026'):
            self.assertContains(r, esperado)
        mes = r.context['d']['mes']
        self.assertEqual((mes['utilidad'], mes['dias_equilibrio'], mes['dias_vendidos'], mes['falta_cf']),
                         (D('16250.00'), 11, 18, D('0')))

    def test_la_meta_cae_al_plan_del_mes(self):
        r = self.client.get(reverse('indicadores:tablero'), {'mes': '2026-02'})
        mes = r.context['d']['mes']
        plan = r.context['d']['plan']
        self.assertEqual(mes['fuente_meta'], 'plan')
        self.assertEqual(mes['meta_venta'], plan['fila']['venta'])
        self.assertNotIn('mezcla', r.context['d'])

    def test_una_meta_capturada_manda_sobre_el_plan(self):
        from indicadores.models import MetaMensual
        MetaMensual.objects.create(anio=2026, mes=2, meta_venta_neta=D('90000'), meta_ocupacion_pct=D('50'))
        mes = self.client.get(reverse('indicadores:tablero'), {'mes': '2026-02'}).context['d']['mes']
        self.assertEqual(mes['fuente_meta'], 'meta')
        self.assertEqual(mes['meta_venta'], D('90000.00'))

    def test_la_simulacion_por_get_ya_no_cambia_nada(self):
        limpio = self.client.get(reverse('indicadores:tablero'), {'mes': '2026-02'}).context['d']['mes']
        r = self.client.get(reverse('indicadores:tablero'), {'mes': '2026-02', 'ticket_EVENTO': '20000'})
        self.assertEqual(r.context['d']['mes'], limpio)
        self.assertNotContains(r, 'Simulación activa')

    def test_sin_clases_de_acento_ni_estilos_en_linea(self):
        for url, params in ((reverse('indicadores:tablero'), {'mes': '2026-02'}), (reverse('indicadores:plan'), {})):
            html = self.client.get(url, params).content.decode()
            self.assertNotIn('qkt-panel--', html, msg=url)
            self.assertEqual(re.findall(r'\sstyle="(?!width:)', html), [], msg=url)

    def test_el_plan_muestra_que_vender_y_mes_por_mes(self):
        r = self.client.get(reverse('indicadores:plan'), {'anio': 2026})
        self.assertContains(r, 'Qué vender')
        self.assertContains(r, 'Mes por mes')
        self.assertNotContains(r, 'Si logras')


class SinPreciosTest(TestCase):

    def test_sin_precios_ni_meta_la_meta_cae_al_equilibrio(self):
        login_superuser_con_totp(self.client, User.objects.create_superuser('dir', 'd@x.mx', 'clave-de-prueba-123'))
        mes = self.client.get(reverse('indicadores:tablero'), {'mes': '2026-02'}).context['d']['mes']
        self.assertIn(mes['fuente_meta'], ('equilibrio', None))
        self.assertNotEqual(mes['fuente_meta'], 'meta')


class SistemaDeDisenoTest(TestCase):

    def test_plantillas_sin_style_ni_emojis_ni_confirm_nativo(self):
        raiz = Path(settings.BASE_DIR) / 'indicadores' / 'templates'
        for ruta in raiz.rglob('*.html'):
            texto = ruta.read_text()
            if 'documentos/_base.html' in texto:
                continue                                  # los PDF tienen su propia guardia
            self.assertNotRegex(texto, r'\sstyle="(?!width:)', msg=ruta.name)   # solo el ancho de la barra de avance
            self.assertIsNone(EMOJI.search(texto), msg=ruta.name)
            self.assertNotIn('confirm(', texto, msg=ruta.name)
            for version in re.findall(r'Chart\.js/([\w.]+)/', texto):
                self.assertNotEqual(version, 'latest')




class EnlacesTest(TestCase):

    def test_el_centro_de_reportes_enlaza_solo_a_direccion(self):
        direccion = User.objects.create_superuser('direccion', 'd@x.mx', 'clave-de-prueba-123')
        staff = User.objects.create_user('staff', 's@x.mx', 'clave-de-prueba-123', is_staff=True)
        url = reverse('indicadores:tablero')
        login_superuser_con_totp(self.client, direccion)
        self.assertContains(self.client.get(reverse('reportes:selector')), url)
        self.client.logout()
        self.client.force_login(staff)
        self.assertNotContains(self.client.get(reverse('reportes:selector')), url)


class ExportacionTest(MesArmado):

    def setUp(self):
        self.precios_de_febrero(guardar=True)
        login_superuser_con_totp(self.client, User.objects.create_superuser('dir', 'd@x.mx', 'clave-de-prueba-123'))

    def test_pdf(self):
        import io

        import pdfplumber
        r = self.client.get(reverse('indicadores:exportar'), {'mes': '2026-02'})
        self.assertEqual(r['Content-Type'], 'application/pdf')
        self.assertIn('QKT_Indicadores_2026-02.pdf', r['Content-Disposition'])
        with pdfplumber.open(io.BytesIO(r.content)) as pdf:
            texto = '\n'.join(p.extract_text() or '' for p in pdf.pages)
        # Plan de febrero a mano: 8 eventos (2 por semana, sáb-dom) y 20 pasadías; 8 × 11,000 + 20 × 3,000 = 148,000
        for esperado in ('Indicadores del mes', '$50,000.00', '$148,000.00', '$63,700.00', '38.6 %', 'Plan óptimo · 02/2026'):
            self.assertIn(esperado, texto)

    def test_excel(self):
        import io

        import openpyxl
        r = self.client.get(reverse('indicadores:exportar'), {'mes': '2026-02', 'formato': 'excel'})
        self.assertIn('QKT_Indicadores_2026-02.xlsx', r['Content-Disposition'])
        libro = openpyxl.load_workbook(io.BytesIO(r.content))
        self.assertIn('Plan anual', libro.sheetnames)
        self.assertNotIn('Escenarios', libro.sheetnames)
        filas = list(libro['Plan anual'].iter_rows(min_row=2, values_only=True))
        self.assertEqual(len(filas), 13)                                  # 12 meses y el año
        self.assertEqual(sum(f[5] for f in filas[:12]), filas[12][5])    # la venta mensual suma la anual
        equilibrio = {f[0].value: f[1].value for f in libro['Equilibrio'].iter_rows(min_row=2)}
        self.assertEqual(equilibrio['Venta neta necesaria'], 50000)

    def test_staff_no_exporta(self):
        self.client.logout()
        self.client.force_login(User.objects.create_user('s', 's@x.mx', 'clave-de-prueba-123', is_staff=True))
        self.assertEqual(self.client.get(reverse('indicadores:exportar')).status_code, 403)


class ProponerMetaTest(MesArmado):

    def setUp(self):
        self.precios_de_febrero(guardar=True)

    def test_con_precios_la_meta_sale_del_plan(self):
        from indicadores.services.metas import proponer_meta
        from indicadores.services.plan_anual import PlanAnualService

        meta = proponer_meta(2026, 2, hoy=date(2026, 3, 5), parametros=self.p)
        fila = PlanAnualService.del_mes(2026, 2, self.p, date(2026, 3, 5))['fila']
        self.assertEqual((meta.meta_eventos, meta.meta_pasadias, meta.meta_noches, meta.meta_venta_neta),
                         (fila['eventos'], fila['pasadias'], fila['noches'], fila['venta']))
        self.assertIn('Propuesta desde plan', meta.notas)

    def test_propone_y_actualiza_sin_duplicar(self):
        from unittest import mock

        from indicadores.models import MetaMensual
        from indicadores.services.metas import proponer_meta

        usuario = User.objects.create_superuser('dir', 'd@x.mx', 'clave-de-prueba-123')
        with mock.patch('indicadores.services.metas.valores_desde_plan', return_value=None):   # sin plan: cascada
            meta = proponer_meta(2026, 2, usuario, hoy=date(2026, 3, 5), parametros=self.p)
        self.assertEqual((meta.meta_venta_neta, meta.meta_ocupacion_pct), (D('50000.00'), D('38.6')))
        self.assertEqual((meta.meta_eventos, meta.meta_pasadias, meta.meta_noches), (3, 5, 12))   # 9 ventas repartidas 4:6:2 (half-up); 2 reservaciones × 6 habitaciones-noche
        self.assertEqual(meta.presupuesto_marketing, D('4431.75'))
        self.assertEqual(meta.created_by, usuario)
        with mock.patch('indicadores.services.metas.valores_desde_plan', return_value=None):
            meta_otra = proponer_meta(2026, 2, usuario, hoy=date(2026, 3, 5), parametros=self.p)
        self.assertEqual(meta.pk, meta_otra.pk)
        self.assertEqual(MetaMensual.objects.filter(anio=2026, mes=2, activo=True).count(), 1)

    def test_accion_del_admin(self):
        from django.urls import reverse as r

        from indicadores.models import MetaMensual
        usuario = User.objects.create_superuser('dir', 'd@x.mx', 'clave-de-prueba-123')
        login_superuser_con_totp(self.client, usuario)
        meta = MetaMensual.objects.create(anio=2026, mes=2)
        resp = self.client.post(r('admin:indicadores_metamensual_changelist'), {
            'action': 'proponer_desde_equilibrio', '_selected_action': [meta.pk]})
        self.assertEqual(resp.status_code, 302)
        meta.refresh_from_db()
        self.assertGreater(meta.meta_venta_neta, 0)
